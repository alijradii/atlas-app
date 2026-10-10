"""Real extraction fixtures and hostile-input checks; no provider requests."""

import hashlib
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from rag_service.documents import (
    FORMATS,
    identify_upload,
    parser_requirements,
    read_document_text,
    validate_upload,
)
from rag_service.parse_document import SafeHTMLBlocks, parse_document


def source_folder(tmp_path, filename, content):
    folder = tmp_path / Path(filename).stem
    folder.mkdir()
    info = identify_upload(filename)
    path = folder / ("original" + info.extension)
    path.write_bytes(content)
    (folder / "source.json").write_text(
        json.dumps(
            {
                "filename": filename,
                "stored_name": path.name,
                "format": info.format,
                "media_type": info.media_type,
                "source_sha256": hashlib.sha256(content).hexdigest(),
            }
        )
    )
    return folder


def content_rows(folder):
    return json.loads((folder / "content.json").read_text())


@pytest.mark.parametrize("extension", FORMATS)
def test_format_registry_accepts_uppercase_and_binary_mime(extension):
    info = identify_upload("UPLOAD" + extension.upper(), "application/octet-stream")
    assert info.extension == extension
    assert identify_upload("file" + extension, info.media_type + "; charset=utf-8") == info


def test_unknown_formats_and_conflicting_mime_fail():
    with pytest.raises(ValueError, match="unsupported_document_format"):
        identify_upload("audio.mp3", "audio/mpeg")
    with pytest.raises(ValueError, match="content_type_mismatch"):
        identify_upload("notes.md", "application/pdf")


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16", "utf-32"])
def test_strict_text_decodes_multilingual_bom_encodings(tmp_path, encoding):
    path = tmp_path / "source.txt"
    expected = "مرحبا بالعالم\nPressure is ٣ units."
    path.write_bytes(expected.encode(encoding))
    assert read_document_text(path) == expected


@pytest.mark.parametrize("raw", [b"binary\x00payload", b"\xffinvalid", b"  \n"])
def test_invalid_text_is_rejected_without_replacement(tmp_path, raw):
    path = tmp_path / "source.txt"
    path.write_bytes(raw)
    with pytest.raises(ValueError):
        validate_upload(path, identify_upload("source.txt"))


def test_pdf_page_count_and_encryption(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "source.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    assert validate_upload(path, FORMATS[".pdf"]) == {"pages": 1, "page_basis": "physical"}
    writer.encrypt("secret")
    writer.write(path)
    with pytest.raises(ValueError, match="encrypted_pdf"):
        validate_upload(path, FORMATS[".pdf"])


def test_office_zip_rejects_traversal_wrong_family_and_expansion(tmp_path):
    path = tmp_path / "upload.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "document")
        archive.writestr("../escape", "no")
    with pytest.raises(ValueError, match="archive_path"):
        validate_upload(path, FORMATS[".docx"])
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("ppt/presentation.xml", "presentation")
    with pytest.raises(ValueError, match="document_structure"):
        validate_upload(path, FORMATS[".docx"])
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", b"0" * 2_000_000)
    with pytest.raises(ValueError, match="compression_limit"):
        validate_upload(path, FORMATS[".docx"])


def test_legacy_converter_requirement_is_actionable(tmp_path, monkeypatch):
    monkeypatch.setattr("rag_service.documents.shutil.which", lambda name: None)
    requirements = parser_requirements(FORMATS[".doc"], tmp_path / "missing.json")
    assert any("LibreOffice" in item for item in requirements)
    assert any("MINERU_TOOLS_CONFIG_JSON" in item for item in requirements)
    assert parser_requirements(FORMATS[".html"], None) == []


def test_markdown_preserves_sections_code_and_literal_image_references(tmp_path):
    pytest.importorskip("raganything")
    from PIL import Image

    secret_image = tmp_path / "outside.png"
    Image.new("RGB", (4, 4), "red").save(secret_image)
    text = (
        "# Safety\n\nPressure limit is 12.\n\n## Procedure\n\n- Close valve\n- Check meter\n\n"
        "```python\n# this is code\nprint('ready')\n```\n\n"
        f"![private]({secret_image})\n![remote](https://example.invalid/asset.png)\n"
    )
    folder = source_folder(tmp_path, "guide.md", text.encode())
    result = parse_document(folder)
    rows = content_rows(folder)
    assert [(row["text"], row["text_level"]) for row in rows if "text_level" in row] == [
        ("Safety", 1),
        ("Procedure", 2),
    ]
    assert all(row["type"] == "text" for row in rows)
    assert any("# this is code" in row["text"] and "```python" in row["text"] for row in rows)
    assert any(str(secret_image) in row["text"] for row in rows)
    assert result["page_basis"] == "logical"
    assert result["parser"] == "raganything-text"
    assert result["source_sha256"] == hashlib.sha256(text.encode()).hexdigest()


def test_plain_text_retains_long_tail(tmp_path):
    pytest.importorskip("raganything")
    text = "A long paragraph. " * 5000 + "TAIL MUST SURVIVE.\n\nSecond paragraph."
    folder = source_folder(tmp_path, "notes.txt", text.encode())
    parse_document(folder)
    rows = content_rows(folder)
    assert len(rows) == 2
    assert rows[0]["text"].endswith("TAIL MUST SURVIVE.")
    assert sum(len(row["text"]) for row in rows) == len(text) - 2


def test_html_preserves_structure_without_active_or_remote_content(tmp_path):
    document = """<!doctype html><html><head><title>Hidden</title></head><body>
    <h1>Pressure</h1><p>A <strong>safe</strong> &amp; clear paragraph.</p>
    <script>secretScript()</script><style>secretStyle</style>
    <h2>Steps</h2><ul><li>Close valve</li><li>Check meter</li></ul>
    <pre>if ready:\n    open_valve()</pre>
    <table><tr><th>Unit</th><th>Value</th></tr><tr><td>Pa</td><td>12</td></tr></table>
    <img src="file:///etc/passwd" alt="pressure diagram"><iframe src="https://example.invalid"></iframe>
    </body></html>"""
    folder = source_folder(tmp_path, "guide.html", document.encode())
    metadata = parse_document(folder)
    rows = content_rows(folder)
    body = json.dumps(rows)
    assert "secretScript" not in body and "secretStyle" not in body and "Hidden" not in body
    assert any(row.get("text_level") == 2 and row["text"] == "Steps" for row in rows)
    assert any(
        row.get("block_kind") == "code" and "\n    open_valve()" in row["text"] for row in rows
    )
    assert any(row["type"] == "table" and "<td>12</td>" in row["table_body"] for row in rows)
    assert "file:///etc/passwd" not in body and "https://example.invalid" not in body
    assert metadata["parser"] == "safe-html"


def test_malformed_html_keeps_trailing_text():
    parser = SafeHTMLBlocks()
    parser.feed("<h1>Heading</h1><p>Unclosed paragraph")
    assert parser.finish()[-1]["text"] == "Unclosed paragraph"


def test_parser_cli_works_outside_service_directory(tmp_path):
    import rag_service.parse_document as module

    folder = source_folder(tmp_path, "cli.html", b"<h1>CLI input</h1><p>Preserved text.</p>")
    result = subprocess.run(
        [sys.executable, module.__file__, str(folder)],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        timeout=20,
    )
    assert result.returncode == 0, result.stderr
    assert content_rows(folder)[0]["text"] == "CLI input"


def test_html_recovers_open_cells_and_flattens_nested_table_text():
    parser = SafeHTMLBlocks()
    parser.feed("<table><tr><td>Outer<table><tr><td>Inner</td></tr></table>Tail<td>Final")
    body = parser.finish()[0]["table_body"]
    assert "Outer Inner Tail" in body
    assert "<td>Final</td>" in body


def test_image_frame_limits_count_each_frame_dimensions(tmp_path, monkeypatch):
    from PIL import Image

    path = tmp_path / "frames.tiff"
    Image.new("RGB", (2, 2)).save(path, save_all=True, append_images=[Image.new("RGB", (10, 10))])
    assert validate_upload(path, FORMATS[".tiff"])["pages"] == 2
    monkeypatch.setattr("rag_service.documents.MAX_IMAGE_PIXELS", 100)
    with pytest.raises(ValueError, match="image_pixel_or_frame_limit"):
        validate_upload(path, FORMATS[".tiff"])


def test_multiframe_raster_extracts_all_frames(tmp_path, monkeypatch):
    pytest.importorskip("raganything")
    from PIL import Image

    from rag_service.parse_document import _parse_mineru

    path = tmp_path / "frames.tiff"
    Image.new("RGB", (8, 8)).save(
        path, save_all=True, append_images=[Image.new("RGB", (8, 8), "white")]
    )
    visited = []

    class FakeParser:
        def parse_image(self, path, **kwargs):
            visited.append(path)
            return [{"type": "text", "text": f"Frame {len(visited)}", "page_idx": 0}]

    monkeypatch.setattr("raganything.parser.get_parser", lambda name: FakeParser())
    rows, preview = _parse_mineru(path, tmp_path, "image")
    assert len(visited) == 2
    assert [row["page_idx"] for row in rows] == [0, 1]
    assert preview is None


@pytest.mark.parametrize("extension", [".png", ".jpg", ".bmp", ".tif", ".gif", ".webp"])
def test_real_image_content_validation(tmp_path, extension):
    from PIL import Image

    path = tmp_path / ("sample" + extension)
    Image.new("RGB", (8, 8), "white").save(path)
    assert validate_upload(path, FORMATS[extension])["pages"] == 1
    if extension != ".png":
        with pytest.raises(ValueError, match="image_content_type_mismatch"):
            validate_upload(path, FORMATS[".png"])


def test_real_docx_native_extraction_preserves_headings_lists_and_table(tmp_path):
    pytest.importorskip("mineru")
    from docx import Document

    path = tmp_path / "native.docx"
    doc = Document()
    doc.add_heading("Pressure manual", 1)
    doc.add_paragraph("Valve opens at twelve units.")
    doc.add_heading("Procedure", 2)
    doc.add_paragraph("Close the valve", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Unit"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Pa"
    table.cell(1, 1).text = "12"
    doc.save(path)
    folder = source_folder(tmp_path, "manual.docx", path.read_bytes())
    metadata = parse_document(folder)
    rows = content_rows(folder)
    assert any(row.get("text_level") == 1 and "Pressure manual" in row["text"] for row in rows)
    assert any("Valve opens" in row.get("text", "") for row in rows)
    assert any("Close the valve" in row.get("text", "") for row in rows)
    assert any(row["type"] == "table" and "12" in row["table_body"] for row in rows)
    assert metadata["parser"] == "mineru-office"
    assert metadata["page_basis"] == "logical"


def test_real_pptx_and_xlsx_native_extraction(tmp_path):
    pytest.importorskip("mineru")
    from openpyxl import Workbook
    from pptx import Presentation

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Pressure lesson"
    slide.placeholders[1].text = "Check the gauge before opening."
    pptx = tmp_path / "native.pptx"
    slides.save(pptx)
    folder = source_folder(tmp_path, "lesson.pptx", pptx.read_bytes())
    parse_document(folder)
    assert "Check the gauge" in json.dumps(content_rows(folder))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet.append(["Pressure", "Volume"])
    sheet.append([12, 24])
    xlsx = tmp_path / "native.xlsx"
    workbook.save(xlsx)
    folder = source_folder(tmp_path, "measurements.xlsx", xlsx.read_bytes())
    parse_document(folder)
    body = json.dumps(content_rows(folder))
    assert "Pressure" in body and "24" in body
